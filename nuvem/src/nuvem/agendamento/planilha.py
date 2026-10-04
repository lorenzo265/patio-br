"""Importação de planilha (SDD 3.4): o gestor sobe a agenda em CSV ou XLSX.

- **Colunas** do modelo, em qualquer ordem, sem ligar para maiúsculas e acentos; colunas a mais
  são ignoradas, e alguns nomes comuns valem pelo do modelo ("data" por "dia").
- **Valores:** dia em ``dd/mm/aaaa`` (ou a data do próprio XLSX) e horas em ``hh:mm``, no fuso
  do site. A planilha não confere o horário de operação nem se a janela passou: é o dado do
  próprio cliente, corrigido e reimportado ao longo do dia.
- **Arquivo:** CSV em UTF-8 ou Windows-1252 (o Excel no Brasil), separado por ``;``, ``,`` ou
  tabulação; ou a primeira aba do XLSX. Até 5 MB e 5.000 linhas. O XLSX é lido com o defusedxml
  (contra XML feito para atacar o leitor) e só se o conteúdo descompactado couber em 20 MB.
- **Relatório por linha:** cada linha certa vira ``Lido``; cada errada, ``Recusado`` com o número
  da linha. Problema no arquivo inteiro levanta ``PlanilhaInvalidaError`` antes de gravar nada.
"""

import csv
import io
import re
import unicodedata
import zipfile
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import UTC, date, datetime, time, timedelta
from zoneinfo import ZoneInfo

import openpyxl
from sqlalchemy.orm import Session

from nuvem.agendamento import servico
from nuvem.agendamento.conectores import Lido, Recusado, ler_dados
from nuvem.agendamento.modelos import Origem
from nuvem.agendamento.servico import Relatorio
from nuvem.cadastro import servico as cadastro
from nuvem.cadastro.acesso import Acesso
from nuvem.erros import DadoInvalidoError

MAXIMO_DE_BYTES = 5 * 1024 * 1024
MAXIMO_DESCOMPACTADO = 20 * 1024 * 1024
"""O XLSX é um zip: um arquivo pequeno pode esconder um conteúdo enorme."""
MAXIMO_DE_LINHAS = 5000


@dataclass(frozen=True)
class Coluna:
    """Uma coluna do modelo: o nome que aparece e os nomes que valem por ele."""

    campo: str
    nome: str
    outros_nomes: tuple[str, ...] = ()
    obrigatoria: bool = False


COLUNAS = (
    Coluna("codigo", "código", ("codigo do agendamento", "agendamento"), obrigatoria=True),
    Coluna("dia", "dia", ("data",), obrigatoria=True),
    Coluna(
        "inicio", "início", ("hora inicio", "hora de inicio", "inicio da janela"), obrigatoria=True
    ),
    Coluna("fim", "fim", ("hora fim", "hora de fim", "fim da janela"), obrigatoria=True),
    Coluna("tipo", "tipo", ("operacao",), obrigatoria=True),
    Coluna(
        "placa_cavalo", "placa do cavalo", ("placa", "cavalo", "placa cavalo"), obrigatoria=True
    ),
    Coluna("reboque_1", "reboque 1", ("reboque", "placa do reboque 1", "placa reboque 1")),
    Coluna("reboque_2", "reboque 2", ("placa do reboque 2", "placa reboque 2")),
    Coluna("reboque_3", "reboque 3", ("placa do reboque 3", "placa reboque 3")),
    Coluna("motorista", "motorista", ("nome do motorista",)),
    Coluna("celular", "celular", ("celular do motorista", "telefone")),
    Coluna("toneladas", "toneladas", ("tonelada",)),
    Coluna("chave_nfe", "chave da NF-e", ("nf e", "nfe", "chave nfe", "chave da nfe")),
)
"""As colunas do modelo, na ordem em que ele as traz."""

EXEMPLO = (
    "AG-1", "05/11/2026", "08:00", "10:00", "descarga", "ABC1D23", "DEF4G56", "", "",
    "Motorista Inventado", "(11) 98765-4321", "32,5", "",
)  # fmt: skip
"""Uma linha inventada, na aba de exemplo do modelo XLSX."""


class PlanilhaInvalidaError(DadoInvalidoError):
    """O arquivo inteiro não serve (formato, tamanho, colunas): nada é gravado."""


@dataclass(frozen=True)
class Planilha:
    """O arquivo que o gestor subiu."""

    nome: str
    conteudo: bytes


Celula = str | int | float | bool | date | time | None


class ConectorDaPlanilha:
    """O conector da planilha (SDD 3.4): cada linha vira um agendamento no formato interno."""

    def __init__(self, fuso: str) -> None:
        """Prepara o conector para o site (as horas da planilha são no fuso dele)."""
        self._fuso = ZoneInfo(fuso)

    @property
    def origem(self) -> Origem:
        """Os agendamentos da planilha têm origem ``planilha``."""
        return "planilha"

    def ler(self, entrada: Planilha) -> list[Lido | Recusado]:
        """Cada linha da planilha, entendida ou recusada.

        Raises:
            PlanilhaInvalidaError: se o arquivo inteiro não servir (formato, tamanho, colunas
                obrigatórias, linhas demais).
        """
        linhas = _linhas_do_arquivo(entrada)
        if not linhas:
            raise PlanilhaInvalidaError("a planilha está vazia")
        posicoes = _colunas(linhas[0])
        com_dados = [
            (numero, linha)
            for numero, linha in enumerate(linhas[1:], start=2)
            if any(_texto(celula) for celula in linha)
        ]
        if len(com_dados) > MAXIMO_DE_LINHAS:
            raise PlanilhaInvalidaError(f"a planilha pode ter no máximo {MAXIMO_DE_LINHAS} linhas")
        return [self._ler_linha(numero, linha, posicoes) for numero, linha in com_dados]

    def _ler_linha(
        self, numero: int, linha: Sequence[Celula], posicoes: dict[str, int]
    ) -> Lido | Recusado:
        def valor(campo: str) -> Celula:
            posicao = posicoes.get(campo)
            return linha[posicao] if posicao is not None and posicao < len(linha) else None

        erros: list[str] = []
        dia = _dia(valor("dia"), erros)
        inicio = _hora(valor("inicio"), "início", erros)
        fim = _hora(valor("fim"), "fim", erros)
        if dia is not None and inicio is not None and fim is not None:
            janela = (
                datetime.combine(dia, inicio, self._fuso),
                datetime.combine(dia, fim, self._fuso),
            )
        else:
            # Sem janela, os outros campos ainda são conferidos, com uma janela qualquer.
            qualquer = datetime(2000, 1, 1, tzinfo=UTC)
            janela = (qualquer, qualquer + timedelta(hours=1))
        reboques = [_texto(valor(f"reboque_{n}")) for n in (1, 2, 3)]
        campos = {
            "codigo_externo": _texto(valor("codigo")),
            "janela_inicio": janela[0],
            "janela_fim": janela[1],
            "tipo": _sem_acento(_texto(valor("tipo"))).lower(),
            "placa_cavalo": _texto(valor("placa_cavalo")),
            "placas_reboques": [placa for placa in reboques if placa],
            "motorista_nome": _texto(valor("motorista")),
            "motorista_celular": _texto(valor("celular")),
            "toneladas": _texto(valor("toneladas")),
            "chave_nfe": _texto(valor("chave_nfe")),
        }
        lido = ler_dados(f"linha {numero}", campos)
        if not erros:
            return lido
        if isinstance(lido, Recusado):
            erros.append(lido.motivo)
        return Recusado(onde=f"linha {numero}", motivo="; ".join(erros))


def importar_planilha(
    sessao: Session, acesso: Acesso, site_id: int, planilha: Planilha, *, agora: datetime
) -> Relatorio:
    """Importa a planilha num site do gestor e devolve o relatório por linha.

    Raises:
        NaoEncontradoError: se o site não existir ou não for visível para este usuário.
        PlanilhaInvalidaError: se o arquivo inteiro não servir; nada é gravado.
    """
    site = cadastro.obter_site(sessao, acesso, site_id)
    destino = servico.site_para_agendar(sessao, acesso, site.id)
    conector = ConectorDaPlanilha(fuso=site.fuso)
    return servico.importar(sessao, destino, conector, planilha, agora=agora)


# --- O modelo para baixar ---------------------------------------------------------------------


def modelo_csv() -> bytes:
    """O modelo em CSV, só com o cabeçalho, como o Excel no Brasil abre (``;`` e BOM)."""
    return ("﻿" + ";".join(coluna.nome for coluna in COLUNAS) + "\n").encode()


def modelo_xlsx() -> bytes:
    """O modelo em XLSX: a aba da agenda, só com o cabeçalho, e uma aba com um exemplo."""
    pasta = openpyxl.Workbook()
    agenda = pasta.active
    assert agenda is not None
    agenda.title = "agenda"
    agenda.append([coluna.nome for coluna in COLUNAS])
    exemplo = pasta.create_sheet("exemplo")
    exemplo.append([coluna.nome for coluna in COLUNAS])
    exemplo.append(list(EXEMPLO))
    saida = io.BytesIO()
    pasta.save(saida)
    return saida.getvalue()


# --- Ler o arquivo ----------------------------------------------------------------------------


def _linhas_do_arquivo(planilha: Planilha) -> list[Sequence[Celula]]:
    if len(planilha.conteudo) > MAXIMO_DE_BYTES:
        raise PlanilhaInvalidaError("o arquivo é grande demais (até 5 MB)")
    extensao = planilha.nome.lower().rsplit(".", 1)[-1] if "." in planilha.nome else ""
    if extensao == "csv":
        return _linhas_do_csv(planilha.conteudo)
    if extensao == "xlsx":
        return _linhas_do_xlsx(planilha.conteudo)
    raise PlanilhaInvalidaError("a planilha precisa ser .csv ou .xlsx")


def _linhas_do_csv(conteudo: bytes) -> list[Sequence[Celula]]:
    try:
        texto = conteudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = conteudo.decode("cp1252", errors="replace")  # o Excel no Brasil
    primeira = texto.split("\n", 1)[0]
    separador = max((";", ",", "\t"), key=primeira.count)
    return list(csv.reader(io.StringIO(texto), delimiter=separador))


def _linhas_do_xlsx(conteudo: bytes) -> list[Sequence[Celula]]:
    try:
        with zipfile.ZipFile(io.BytesIO(conteudo)) as compactado:
            descompactado = sum(arquivo.file_size for arquivo in compactado.infolist())
    except zipfile.BadZipFile as erro:
        raise PlanilhaInvalidaError("o arquivo .xlsx não abre") from erro
    if descompactado > MAXIMO_DESCOMPACTADO:
        raise PlanilhaInvalidaError("o conteúdo do .xlsx é grande demais")
    try:
        pasta = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except Exception as erro:
        # O openpyxl tem muitos jeitos de recusar um arquivo estragado; para quem subiu, é um só.
        raise PlanilhaInvalidaError("o arquivo .xlsx não abre") from erro
    try:
        aba = pasta.worksheets[0]
        linhas: list[Sequence[Celula]] = []
        for linha in aba.iter_rows(values_only=True):
            linhas.append(linha)  # type: ignore[arg-type]
            if len(linhas) > MAXIMO_DE_LINHAS + 1:
                break
        return linhas
    finally:
        pasta.close()


def _colunas(cabecalho: Sequence[Celula]) -> dict[str, int]:
    por_nome = {
        _normalizar(nome): coluna
        for coluna in COLUNAS
        for nome in (coluna.nome, *coluna.outros_nomes)
    }
    posicoes: dict[str, int] = {}
    for posicao, celula in enumerate(cabecalho):
        coluna = por_nome.get(_normalizar(_texto(celula)))
        if coluna is None:
            continue
        if coluna.campo in posicoes:
            raise PlanilhaInvalidaError(f"a coluna {coluna.nome} aparece duas vezes")
        posicoes[coluna.campo] = posicao
    faltam = [c.nome for c in COLUNAS if c.obrigatoria and c.campo not in posicoes]
    if faltam:
        raise PlanilhaInvalidaError(f"faltam as colunas: {', '.join(faltam)}")
    return posicoes


# --- Ler cada célula --------------------------------------------------------------------------

_HORA_FALADA = re.compile(r"^(\d{1,2})h(\d{2})?$")


def _dia(celula: Celula, erros: list[str]) -> date | None:
    if isinstance(celula, datetime):
        return celula.date()
    if isinstance(celula, date):
        return celula
    texto = _texto(celula)
    if not texto:
        erros.append("dia: falta")
        return None
    for formato in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    erros.append("dia: use dd/mm/aaaa")
    return None


def _hora(celula: Celula, rotulo: str, erros: list[str]) -> time | None:
    if isinstance(celula, datetime):
        return celula.time()
    if isinstance(celula, time):
        return celula
    texto = _texto(celula).lower()
    if not texto:
        erros.append(f"{rotulo}: falta")
        return None
    falada = _HORA_FALADA.fullmatch(texto)
    try:
        if falada:
            return time(int(falada[1]), int(falada[2] or 0))
        return time.fromisoformat(texto if len(texto) > 4 else f"0{texto}")
    except ValueError:
        erros.append(f"{rotulo}: use hh:mm")
        return None


def _texto(celula: Celula) -> str:
    if celula is None:
        return ""
    if isinstance(celula, float) and celula.is_integer():
        return str(int(celula))  # o Excel guarda 12345 como 12345.0
    return str(celula).strip()


def _sem_acento(texto: str) -> str:
    decomposto = unicodedata.normalize("NFKD", texto)
    return "".join(letra for letra in decomposto if not unicodedata.combining(letra))


def _normalizar(nome: str) -> str:
    return " ".join(re.split(r"[\s_.\-]+", _sem_acento(nome).lower())).strip()
