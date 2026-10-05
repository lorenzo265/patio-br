"""Importação de planilha (SDD 3.4): CSV e XLSX, o modelo, e o relatório por linha."""

import io
import zipfile
from collections.abc import Sequence
from datetime import UTC, date, datetime, time
from decimal import Decimal

import openpyxl
import pytest

from nuvem.agendamento.conectores import Lido, Recusado
from nuvem.agendamento.planilha import (
    MAXIMO_DE_LINHAS,
    ConectorDaPlanilha,
    Planilha,
    PlanilhaInvalidaError,
    modelo_csv,
    modelo_xlsx,
)

CABECALHO = (
    "código;dia;início;fim;tipo;placa do cavalo;reboque 1;reboque 2;reboque 3;"
    "motorista;celular;toneladas;chave da NF-e"
)
LINHA = "AG-1;05/11/2026;08:00;10:00;descarga;ABC1D23;DEF4G56;;;Motorista Inventado;" \
    "(11) 98765-4321;32,5;"  # fmt: skip

CONECTOR = ConectorDaPlanilha(fuso="America/Sao_Paulo")


def _csv(*linhas: str, codificacao: str = "utf-8") -> Planilha:
    return Planilha(nome="agenda.csv", conteudo="\n".join(linhas).encode(codificacao))


def _xlsx(linhas: Sequence[Sequence[object]], nome: str = "agenda.xlsx") -> Planilha:
    pasta = openpyxl.Workbook()
    aba = pasta.active
    assert aba is not None
    for linha in linhas:
        aba.append(list(linha))
    saida = io.BytesIO()
    pasta.save(saida)
    return Planilha(nome=nome, conteudo=saida.getvalue())


def _lidos(planilha: Planilha) -> list[Lido | Recusado]:
    return list(CONECTOR.ler(planilha))


# --- CSV --------------------------------------------------------------------------------------


def test_csv_do_modelo_vira_agendamento_no_fuso_do_site() -> None:
    [lido] = _lidos(_csv(CABECALHO, LINHA))

    assert isinstance(lido, Lido)
    assert lido.onde == "linha 2"
    dados = lido.dados
    assert dados.codigo_externo == "AG-1"
    assert dados.janela_inicio == datetime(2026, 11, 5, 11, 0, tzinfo=UTC)  # 8h em São Paulo
    assert dados.janela_fim == datetime(2026, 11, 5, 13, 0, tzinfo=UTC)
    assert (dados.tipo, dados.placa_cavalo, dados.placas_reboques) == (
        "descarga",
        "ABC1D23",
        ("DEF4G56",),
    )
    assert (dados.motorista_celular, dados.toneladas) == ("+5511987654321", Decimal("32.5"))


def test_csv_do_excel_no_brasil_em_windows_1252() -> None:
    linha = LINHA.replace("Motorista Inventado", "João Inventado")

    [lido] = _lidos(_csv(CABECALHO, linha, codificacao="cp1252"))

    assert isinstance(lido, Lido)
    assert lido.dados.motorista_nome == "João Inventado"


def test_csv_com_bom_e_virgula() -> None:
    cabecalho = "codigo,dia,inicio,fim,tipo,placa do cavalo"
    texto = f"\ufeff{cabecalho}\nAG-2,2026-11-05,8:00,9:30,carga,abc-1234"
    planilha = Planilha(nome="agenda.csv", conteudo=texto.encode())

    [lido] = _lidos(planilha)

    assert isinstance(lido, Lido)
    assert (lido.dados.codigo_externo, lido.dados.tipo, lido.dados.placa_cavalo) == (
        "AG-2",
        "carga",
        "ABC1234",
    )


def test_cabecalho_sem_acento_em_maiusculas_e_fora_de_ordem() -> None:
    planilha = _csv(
        "PLACA;Tipo;Data;FIM;Inicio;Codigo;Observação",
        "ABC1D23;Descarga;05/11/2026;10:00;08:00;AG-3;qualquer coisa",
    )

    [lido] = _lidos(planilha)

    assert isinstance(lido, Lido)
    assert (lido.dados.codigo_externo, lido.dados.tipo) == ("AG-3", "descarga")


def test_linhas_vazias_sao_puladas_e_a_numeracao_segue_a_do_arquivo() -> None:
    planilha = _csv(CABECALHO, "", ";;;;;;;;;;;;", LINHA.replace("ABC1D23", "ABC12"))

    [recusado] = _lidos(planilha)

    assert recusado == Recusado(
        onde="linha 4",
        motivo="placa do cavalo: placa inválida: 'ABC12' (formatos aceitos: ABC1234 e ABC1D23)",
    )


@pytest.mark.parametrize(
    ("trocar", "por", "motivo"),
    [
        ("05/11/2026", "5 de novembro", "dia: use dd/mm/aaaa"),
        ("05/11/2026", "", "dia: falta"),
        ("08:00", "oito", "início: use hh:mm"),
        ("10:00", "", "fim: falta"),
        ("10:00", "07:00", "o fim da janela precisa ser depois do início"),
        ("descarga", "transbordo", "tipo: use carga ou descarga"),
    ],
)
def test_erros_por_linha_dizem_o_campo(trocar: str, por: str, motivo: str) -> None:
    [recusado] = _lidos(_csv(CABECALHO, LINHA.replace(trocar, por, 1)))

    assert recusado == Recusado(onde="linha 2", motivo=motivo)


def test_erro_na_janela_e_em_outro_campo_aparecem_juntos() -> None:
    linha = LINHA.replace("05/11/2026", "amanhã").replace("ABC1D23", "ABC12")

    [recusado] = _lidos(_csv(CABECALHO, linha))

    assert isinstance(recusado, Recusado)
    assert recusado.motivo.startswith("dia: use dd/mm/aaaa; placa do cavalo: placa inválida")


def test_8h_escrito_como_se_fala() -> None:
    [lido] = _lidos(_csv(CABECALHO, LINHA.replace("08:00", "8h").replace("10:00", "10h30")))

    assert isinstance(lido, Lido)
    assert lido.dados.janela_fim == datetime(2026, 11, 5, 13, 30, tzinfo=UTC)


# --- XLSX -------------------------------------------------------------------------------------


def test_xlsx_com_data_e_hora_do_proprio_excel() -> None:
    cabecalho = ["Código", "Dia", "Início", "Fim", "Tipo", "Placa", "Celular", "Toneladas"]
    linha = [12345, datetime(2026, 11, 5), time(8), time(10), "descarga", "ABC1D23", 11987654321]
    planilha = _xlsx([cabecalho, [*linha, 32.5]])

    [lido] = _lidos(planilha)

    assert isinstance(lido, Lido)
    assert lido.dados.codigo_externo == "12345"  # número inteiro, sem ".0"
    assert lido.dados.janela_inicio == datetime(2026, 11, 5, 11, 0, tzinfo=UTC)
    assert lido.dados.motorista_celular == "+5511987654321"
    assert lido.dados.toneladas == Decimal("32.5")


def test_xlsx_com_texto_nas_celulas() -> None:
    planilha = _xlsx([CABECALHO.split(";"), LINHA.split(";")])

    [lido] = _lidos(planilha)

    assert isinstance(lido, Lido)
    assert lido.dados.codigo_externo == "AG-1"


def test_xlsx_com_data_como_date() -> None:
    planilha = _xlsx(
        [
            ["codigo", "dia", "inicio", "fim", "tipo", "placa"],
            ["AG-1", date(2026, 11, 5), "08:00", "10:00", "carga", "ABC1D23"],
        ]
    )

    [lido] = _lidos(planilha)

    assert isinstance(lido, Lido)


# --- O arquivo inteiro ------------------------------------------------------------------------


def test_faltam_colunas_obrigatorias() -> None:
    with pytest.raises(PlanilhaInvalidaError, match="faltam as colunas: dia, início"):
        _lidos(_csv("código;fim;tipo;placa do cavalo", "AG-1;10:00;carga;ABC1D23"))


def test_coluna_repetida() -> None:
    with pytest.raises(PlanilhaInvalidaError, match="a coluna dia aparece duas vezes"):
        _lidos(_csv(CABECALHO + ";data", LINHA + ";05/11/2026"))


def test_arquivo_vazio() -> None:
    with pytest.raises(PlanilhaInvalidaError, match="vazia"):
        _lidos(_csv(""))


@pytest.mark.parametrize("nome", ["agenda.xls", "agenda.ods", "agenda.txt", "agenda"])
def test_so_csv_e_xlsx(nome: str) -> None:
    with pytest.raises(PlanilhaInvalidaError, match=r"\.csv ou \.xlsx"):
        _lidos(Planilha(nome=nome, conteudo=b"qualquer"))


def test_xlsx_que_nao_e_xlsx() -> None:
    with pytest.raises(PlanilhaInvalidaError, match="não abre"):
        _lidos(Planilha(nome="agenda.xlsx", conteudo=b"isto nao e um zip"))


def test_xlsx_que_explode_ao_abrir_e_recusado() -> None:
    # Um zip pequeno com um arquivo enorme por dentro (só zeros, que comprimem muito).
    conteudo = io.BytesIO()
    with zipfile.ZipFile(conteudo, "w", zipfile.ZIP_DEFLATED) as zip_:
        zip_.writestr("xl/worksheets/sheet1.xml", b"\0" * (25 * 1024 * 1024))

    with pytest.raises(PlanilhaInvalidaError, match="grande demais"):
        _lidos(Planilha(nome="agenda.xlsx", conteudo=conteudo.getvalue()))


def test_arquivo_grande_demais() -> None:
    with pytest.raises(PlanilhaInvalidaError, match="grande demais"):
        _lidos(Planilha(nome="agenda.csv", conteudo=b"x" * (5 * 1024 * 1024 + 1)))


def test_linhas_demais() -> None:
    linhas = [CABECALHO] + [LINHA.replace("AG-1", f"AG-{n}") for n in range(MAXIMO_DE_LINHAS + 1)]

    with pytest.raises(PlanilhaInvalidaError, match="no máximo 5000 linhas"):
        _lidos(_csv(*linhas))


# --- O modelo para baixar ---------------------------------------------------------------------


def test_modelo_csv_traz_so_o_cabecalho() -> None:
    # Sem linha de exemplo: o modelo subido sem preencher não cria agendamento nenhum.
    conteudo = modelo_csv()

    assert conteudo.decode("utf-8-sig").splitlines() == [CABECALHO]
    assert _lidos(Planilha(nome="modelo.csv", conteudo=conteudo)) == []


def test_modelo_xlsx_traz_o_cabecalho_e_o_exemplo_em_outra_aba() -> None:
    pasta = openpyxl.load_workbook(io.BytesIO(modelo_xlsx()))
    agenda, exemplo = pasta.worksheets

    assert [celula.value for celula in agenda[1]] == CABECALHO.split(";")
    assert agenda.max_row == 1
    linhas = [[celula.value for celula in linha] for linha in exemplo.iter_rows()]
    assert linhas[0] == CABECALHO.split(";")
    [lido] = _lidos(_xlsx(linhas))
    assert isinstance(lido, Lido)


def test_o_openpyxl_le_com_o_defusedxml() -> None:
    # Sem o defusedxml, o openpyxl aceitaria XML feito para atacar o leitor (SDD 3.4).
    import openpyxl.xml

    assert openpyxl.xml.DEFUSEDXML


def test_modelo_xlsx_subido_sem_preencher_nao_cria_nada() -> None:
    # Só a primeira aba é lida: o exemplo, na segunda, não entra.
    assert _lidos(Planilha(nome="modelo.xlsx", conteudo=modelo_xlsx())) == []


def test_xlsx_de_outro_programa_com_numero_inteiro_escrito_com_ponto() -> None:
    # Há programas que gravam 12345 como "12345.0"; o código continua "12345".
    original = _xlsx([["codigo", "dia", "inicio", "fim", "tipo", "placa"],
                      [12345, "05/11/2026", "08:00", "10:00", "carga", "ABC1D23"]])  # fmt: skip
    entrada, saida = io.BytesIO(original.conteudo), io.BytesIO()
    with zipfile.ZipFile(entrada) as de, zipfile.ZipFile(saida, "w") as para:
        for arquivo in de.infolist():
            dados = de.read(arquivo)
            if arquivo.filename == "xl/worksheets/sheet1.xml":
                assert b"<v>12345</v>" in dados
                dados = dados.replace(b"<v>12345</v>", b"<v>12345.0</v>")
            para.writestr(arquivo, dados)

    [lido] = _lidos(Planilha(nome="agenda.xlsx", conteudo=saida.getvalue()))

    assert isinstance(lido, Lido)
    assert lido.dados.codigo_externo == "12345"
